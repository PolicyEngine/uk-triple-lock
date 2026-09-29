"""Build the OBR forecast-error dataset for the triple lock uncertainty analysis.

Reads raw files in data/raw/ (downloaded 29 Sept 2026, see docs/sources.md) and writes:

- data/obr_forecast_errors.csv               spring/Budget vintages (one per year), horizons 1-5
- data/obr_forecast_errors_all_vintages.csv  every EFO vintage (spring + autumn), horizons 1-5
- data/obr_forecast_error_summary.csv        mean / SD / RMSE by variable and horizon, plus the
                                             CPI-earnings error correlation
- data/obr_central_forecast.csv              March 2026 EFO central forecast (+ CPI fan chart deciles)
- data/obr_outturn_crosscheck.csv            OBR outturn vs ONS D7G7 / AWE (KAB9, KAC3)
- data/triple_lock_actual_inputs.csv         September CPI and May-July AWE total pay growth

Rates are written as decimals (0.021 = 2.1%).
"""

from __future__ import annotations

import csv
import math
import statistics
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
DATA = ROOT / "data"

HOFD_FILE = RAW / "Historical_official_forecasts_database_Spring_2026.xlsx"
HOFD_URL = (
    "https://obr.uk/docs/dlm_uploads/"
    "Historical_official_forecasts_database_Spring_2026.xlsx"
)
EFO_ECON_FILE = RAW / "efo-march-2026-detailed-forecast-tables-economy.xlsx"
EFO_ECON_URL = (
    "https://obr.uk/docs/d055fbf02d5b3g6jq8l2/"
    "efo-march-2026-detailed-forecast-tables-economy.xlsx"
)
EFO_CH2_FILE = RAW / "efo-march-2026-charts-and-tables-chapter-2.xlsx"
EFO_CH2_URL = (
    "https://obr.uk/docs/d055fbf02d5b3g6jq8l2/"
    "efo-march-2026-charts-and-tables-chapter-2.xlsx"
)
LTED_FILE = RAW / "Long-term-economic-determinants-March-2026-EFO.xlsx"
LTED_URL = (
    "https://obr.uk/docs/dlm_uploads/Long-term-economic-determinants-March-2026-EFO.xlsx"
)
ONS_D7G7 = RAW / "ons_d7g7_cpi_annual_rate.csv"
ONS_KAB9 = RAW / "ons_kab9_awe_total_pay.csv"
ONS_KAC3 = RAW / "ons_kac3_awe_total_pay_3m_yoy.csv"
ONS_GEN = "https://www.ons.gov.uk/generator?format=csv&uri="
ONS_EARN_URL = " | ".join(
    ONS_GEN + u
    for u in (
        "/economy/grossdomesticproductgdp/timeseries/dtwm/ukea",
        "/economy/grossdomesticproductgdp/timeseries/royk/ukea",
        "/employmentandlabourmarket/peopleinwork/employmentandemployeetypes/timeseries/mgrz/lms",
        "/employmentandlabourmarket/peopleinwork/employmentandemployeetypes/timeseries/mgrq/lms",
    )
)
ONS_EARN_FILES = {
    "DTWM": RAW / "ons_dtwm_compensation_of_employees.csv",
    "ROYK": RAW / "ons_royk_employers_social_contributions.csv",
    "MGRZ": RAW / "ons_mgrz_employment_16plus.csv",
    "MGRQ": RAW / "ons_mgrq_self_employed_16plus.csv",
}


def ons_quarterly(path):
    out = {}
    with path.open() as f:
        for row in csv.reader(f):
            if len(row) == 2 and " Q" in row[0]:
                try:
                    out[row[0]] = float(row[1])
                except ValueError:
                    pass
    return out


def ons_obr_earnings_growth():
    """Calendar-year growth of OBR 'average earnings' rebuilt from latest ONS data:
    (DTWM - ROYK) / (MGRZ - MGRQ), per the definition in EFO Table 1.6.
    Reproduces the HOFD outturn row to 2dp for 2012-2023 (2024: 5.01 vs 5.09 after revisions)."""
    q = {k: ons_quarterly(v) for k, v in ONS_EARN_FILES.items()}

    def level(y):
        ks = [f"{y} Q{i}" for i in range(1, 5)]
        if not all(all(k in q[s] for k in ks) for s in q):
            return None
        ws = sum(q["DTWM"][k] - q["ROYK"][k] for k in ks)
        emp = sum(q["MGRZ"][k] - q["MGRQ"][k] for k in ks)
        return ws / emp

    out = {}
    for y in range(2000, 2031):
        a, b = level(y), level(y - 1)
        if a and b:
            out[y] = (a / b - 1) * 100
    return out

MAX_HORIZON = 5
MIN_HORIZON = 1
VARIABLES = {"cpi": "CPI", "earnings": "Earnings"}


def read_hofd_sheet(sheet: str):
    """Return ({vintage_label: {year: pct}}, {year: outturn_pct}, [vintage order])."""
    wb = openpyxl.load_workbook(HOFD_FILE, read_only=True, data_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if r[0] == "Back to contents")
    header = rows[header_idx]
    year_cols = {j: int(v) for j, v in enumerate(header) if isinstance(v, (int, float))}
    forecasts, outturn, order = {}, {}, []
    for r in rows[header_idx + 1 :]:
        label = r[0]
        if not isinstance(label, str):
            continue
        values = {
            year_cols[j]: float(r[j])
            for j in year_cols
            if j < len(r) and isinstance(r[j], (int, float))
        }
        if label.startswith("Outturn"):
            outturn = values
            break
        forecasts[label] = values
        order.append(label)
    return forecasts, outturn, order


def vintage_year(label: str) -> int:
    return int(label.split()[-1])


def spring_vintages(order):
    """First forecast of each calendar year: June 2010, then March 2011 onwards."""
    seen, out = set(), []
    for label in order:
        y = vintage_year(label)
        if y not in seen:
            seen.add(y)
            out.append(label)
    return out


def vintage_name(label: str) -> str:
    if label == "June 2010":
        return "June 2010 Budget forecast"
    return f"{label} EFO"


def build_errors(vintage_filter):
    records = []
    for var, sheet in VARIABLES.items():
        forecasts, outturn, order = read_hofd_sheet(sheet)
        outturn_src = {y: HOFD_URL for y in outturn}
        if var == "earnings":
            # HOFD earnings outturn ends in 2024; fill later years from ONS on the OBR definition.
            for y, v in ons_obr_earnings_growth().items():
                if y not in outturn:
                    outturn[y] = v
                    outturn_src[y] = HOFD_URL + " | " + ONS_EARN_URL
        vintages = vintage_filter(order)
        for label in vintages:
            made = vintage_year(label)
            for target, fc in sorted(forecasts[label].items()):
                h = target - made
                if h < MIN_HORIZON or h > MAX_HORIZON:
                    continue
                if target not in outturn:
                    continue
                o = outturn[target]
                records.append(
                    {
                        "year_forecast_made": made,
                        "forecast_vintage": vintage_name(label),
                        "target_year": target,
                        "horizon_years": h,
                        "variable": var,
                        "forecast": round(fc / 100, 6),
                        "outturn": round(o / 100, 6),
                        "error": round((o - fc) / 100, 6),
                        "source_url": outturn_src[target],
                    }
                )
    records.sort(
        key=lambda r: (r["variable"], r["year_forecast_made"], r["forecast_vintage"], r["target_year"])
    )
    return records


def write_csv(path: Path, rows, fields):
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def corr(xs, ys):
    if len(xs) < 3:
        return None
    mx, my = statistics.fmean(xs), statistics.fmean(ys)
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    sxx = sum((x - mx) ** 2 for x in xs)
    syy = sum((y - my) ** 2 for y in ys)
    return sxy / math.sqrt(sxx * syy)


def summarise(records, sample):
    out = []
    by_key = {}
    for r in records:
        by_key.setdefault((r["variable"], r["horizon_years"]), []).append(r)
    # Horizons that occur in this sample; each must have both variables
    # (a KeyError here means the sample is inconsistent, not an empty cell).
    for h in sorted({h for _, h in by_key}):
        cpi = {(r["forecast_vintage"], r["target_year"]): r["error"] for r in by_key[("cpi", h)]}
        ern = {(r["forecast_vintage"], r["target_year"]): r["error"] for r in by_key[("earnings", h)]}
        pairs = sorted(set(cpi) & set(ern))
        rho = corr([cpi[k] for k in pairs], [ern[k] for k in pairs])
        for var in VARIABLES:
            errs = [r["error"] for r in by_key[(var, h)]]
            n = len(errs)
            targets = sorted({r["target_year"] for r in by_key[(var, h)]})
            out.append(
                {
                    "sample": sample,
                    "variable": var,
                    "horizon_years": h,
                    "n": n,
                    "target_years": f"{targets[0]}-{targets[-1]}",
                    "mean_error": round(statistics.fmean(errs), 5),
                    "sd_error": round(statistics.stdev(errs), 5) if n > 1 else "",
                    "rmse": round(math.sqrt(statistics.fmean([e * e for e in errs])), 5),
                    "mean_abs_error": round(statistics.fmean([abs(e) for e in errs]), 5),
                    "corr_cpi_earnings_error": round(rho, 3) if rho is not None else "",
                    "n_pairs": len(pairs),
                }
            )
    return out


# ---------------------------------------------------------------- central forecast


def efo_table(sheet):
    wb = openpyxl.load_workbook(EFO_ECON_FILE, read_only=True, data_only=True)
    return list(wb[sheet].iter_rows(values_only=True))


def efo_col(rows, header_row, name):
    return next(j for j, v in enumerate(rows[header_row]) if isinstance(v, str) and v.strip().startswith(name))


def efo_series(rows, col):
    return {str(r[1]).strip(): float(r[col]) for r in rows if r[1] is not None and isinstance(r[col], (int, float))}


def build_central():
    rows = []

    def add(variable, basis, period, value, source, url, note="", pct=None):
        row = {
            "variable": variable,
            "basis": basis,
            "period": period,
            "value": round(value / 100, 6),
            "p10": "", "p20": "", "p30": "", "p40": "", "p50": "",
            "p60": "", "p70": "", "p80": "", "p90": "",
            "source": source,
            "source_url": url,
            "note": note,
        }
        if pct:
            for k, v in pct.items():
                row[k] = round(v / 100, 6)
        rows.append(row)

    src_tables = "OBR March 2026 EFO, detailed forecast tables: economy"
    # CPI: Table 1.7 (header row index 3)
    t17 = efo_table("1.7")
    cpi = efo_series(t17, efo_col(t17, 3, "CPI"))
    # Earnings: Table 1.6 'Average weekly earnings growth' = wages and salaries / employees
    t16 = efo_table("1.6")
    ern = efo_series(t16, efo_col(t16, 2, "Average weekly earnings growth"))

    # CPI fan chart deciles (Chart 2.9), calendar years
    wb = openpyxl.load_workbook(EFO_CH2_FILE, read_only=True, data_only=True)
    fan = {}
    for r in wb["C2.9"].iter_rows(values_only=True):
        if isinstance(r[1], (int, float)) and isinstance(r[4], (int, float)):
            fan[int(r[1])] = {f"p{p}": float(r[4 + i]) for i, p in enumerate(range(10, 100, 10))}

    for y in range(2026, 2031):
        add("cpi", "calendar_year", str(y), cpi[str(y)], src_tables + " Table 1.7", EFO_ECON_URL,
            "Fan chart deciles from EFO Chart 2.9 (charts and tables, chapter 2)", fan.get(y))
    for y in range(2026, 2031):
        add("cpi", "q3_yoy", f"{y}Q3", cpi[f"{y}Q3"], src_tables + " Table 1.7", EFO_ECON_URL,
            "Q3 average of 12-month CPI rate; triple lock uses September 12-month rate")
    for y in range(2026, 2031):
        add("earnings", "calendar_year", str(y), ern[str(y)], src_tables + " Table 1.6", EFO_ECON_URL,
            "OBR average earnings = wages and salaries / employees (national accounts), not AWE")
    for y in range(2026, 2031):
        add("earnings", "q2_yoy", f"{y}Q2", ern[f"{y}Q2"], src_tables + " Table 1.6", EFO_ECON_URL,
            "Q2 y/y growth of wages and salaries per employee; triple lock uses May-July AWE total pay")
    for y in range(2025, 2031):
        fy = f"{y}-{str(y + 1)[2:]}"
        add("cpi", "fiscal_year", fy, cpi[fy], src_tables + " Table 1.7", EFO_ECON_URL)
        add("earnings", "fiscal_year", fy, ern[fy], src_tables + " Table 1.6", EFO_ECON_URL,
            "wages and salaries / employees")

    # Long-term economic determinants (fiscal years; beyond 2030-31 = long-term projection)
    wb = openpyxl.load_workbook(LTED_FILE, read_only=True, data_only=True)
    lrows = list(wb["Long-term economic determinants"].iter_rows(values_only=True))
    hdr = next(r for r in lrows if r[2] == "2023-24")
    lmap = {"Average earnings": "earnings", "CPI": "cpi", "'Triple Lock'": "triple_lock_uprating"}
    for r in lrows:
        if r[1] in lmap:
            for j in range(2, len(hdr)):
                fy = hdr[j]
                if not isinstance(fy, str) or not isinstance(r[j], (int, float)):
                    continue
                start = int(fy[:4])
                if start < 2025 or start > 2036:
                    continue
                note = "Medium-term EFO forecast" if start <= 2030 else "Long-term projection beyond EFO horizon (not a forecast)"
                if lmap[r[1]] == "triple_lock_uprating":
                    note += ("; OBR 'Triple Lock' row: uprating determined by this fiscal year's inputs "
                             "(Sept CPI / May-Jul AWE), paid from the following April "
                             "(e.g. 2025-26 = 4.8% paid April 2026)")
                if lmap[r[1]] == "earnings":
                    note += "; average earnings = wages and salaries / employees"
                add(lmap[r[1]], "fiscal_year_lted", fy, float(r[j]),
                    "OBR Long-term economic determinants, March 2026 EFO (dated 28 May 2026)", LTED_URL, note)
    return rows


# ---------------------------------------------------------------- ONS cross-checks


def read_ons(path):
    out = {}
    with path.open() as f:
        for row in csv.reader(f):
            if len(row) == 2 and row[0][:4].isdigit():
                try:
                    out[row[0]] = float(row[1])
                except ValueError:
                    pass
    return out


MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


def build_crosscheck():
    d7g7 = read_ons(ONS_D7G7)
    kab9 = read_ons(ONS_KAB9)
    kac3 = read_ons(ONS_KAC3)
    _, cpi_out, _ = read_hofd_sheet("CPI")
    _, ern_out, _ = read_hofd_sheet("Earnings")

    def kab9_avg(y):
        vals = [kab9.get(f"{y} {m}") for m in MONTHS]
        return statistics.fmean(vals) if all(v is not None for v in vals) else None

    ons_e = ons_obr_earnings_growth()
    rows, inputs = [], []
    for y in range(2010, 2027):
        a, b = kab9_avg(y), kab9_avg(y - 1)
        awe_cal = (a / b - 1) * 100 if a and b else None
        e = ern_out.get(y)
        rows.append(
            {
                "year": y,
                "cpi_obr_outturn": _pct(cpi_out.get(y)),
                "cpi_ons_d7g7_calendar": _pct(d7g7.get(str(y))),
                "cpi_ons_d7g7_september": _pct(d7g7.get(f"{y} SEP")),
                "earnings_obr_outturn": _pct(e),
                "earnings_obr_def_rebuilt_latest_ons": _pct(ons_e.get(y)),
                "awe_kab9_calendar_growth": _pct(awe_cal),
                "awe_kac3_may_jul": _pct(kac3.get(f"{y} JUL")),
                "gap_kac3_minus_obr_earnings": _pct(kac3[f"{y} JUL"] - e) if e is not None and f"{y} JUL" in kac3 else "",
            }
        )
        sep, jul = d7g7.get(f"{y} SEP"), kac3.get(f"{y} JUL")
        inputs.append(
            {
                "determination_year": y,
                "uprating_april": y + 1,
                "cpi_september_12m": _pct(sep),
                "awe_total_pay_may_jul_3m_yoy": _pct(jul),
                "max_cpi_awe_2_5pct": _pct(max(sep, jul, 2.5)) if sep is not None and jul is not None else "",
                "note": "Latest ONS vintage (Sept 2026 release); the statutory uprating used first-published figures, which can differ",
            }
        )
    return rows, inputs


def _pct(v):
    return "" if v is None else round(v / 100, 6)


def main():
    spring = build_errors(spring_vintages)
    allv = build_errors(lambda order: order)
    fields = ["year_forecast_made", "forecast_vintage", "target_year", "horizon_years", "variable",
              "forecast", "outturn", "error", "source_url"]
    write_csv(DATA / "obr_forecast_errors.csv", spring, fields)
    write_csv(DATA / "obr_forecast_errors_all_vintages.csv", allv, fields)

    ex_shock = [r for r in spring if not 2020 <= r["target_year"] <= 2023]
    summary = (
        summarise(spring, "spring_vintages")
        + summarise(allv, "all_vintages")
        + summarise(ex_shock, "spring_vintages_excl_targets_2020_2023")
    )
    write_csv(DATA / "obr_forecast_error_summary.csv", summary, list(summary[0].keys()))

    central = build_central()
    write_csv(DATA / "obr_central_forecast.csv", central, list(central[0].keys()))

    cc, inputs = build_crosscheck()
    write_csv(DATA / "obr_outturn_crosscheck.csv", cc, list(cc[0].keys()))
    write_csv(DATA / "triple_lock_actual_inputs.csv", inputs, list(inputs[0].keys()))

    print(f"spring rows: {len(spring)}, all-vintage rows: {len(allv)}, central rows: {len(central)}")


if __name__ == "__main__":
    main()
